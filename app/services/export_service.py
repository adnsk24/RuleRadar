import io
import json
from datetime import datetime, timezone
from typing import Dict, Any
from docx import Document as DocxDoc
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, PageBreak, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

class ExportService:
    """
    Generates professional, printable, and downloadable analysis reports
    in PDF, DOCX, and JSON formats with full citation traceability.
    """

    DISCLAIMER_TEXT = (
        "IMPORTANT LEGAL DISCLAIMER: This report is generated automatically by RuleRadar AI using open-source NLP "
        "and rule-based document intelligence. The risk review indicators, importance ratings, and extracted terms are provided "
        "solely for informational and prioritization purposes and DO NOT constitute legal advice or formal statutory conclusions. "
        "Always review the complete original document and consult a qualified legal professional before signing or acting on legal provisions."
    )

    @classmethod
    def generate_pdf_report(cls, doc_data: Dict[str, Any]) -> bytes:
        """
        Generate a beautifully styled, comprehensive PDF report using ReportLab.
        """
        buffer = io.BytesIO()
        pdf = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()

        # Custom styles
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#1F2937'),
            spaceAfter=6
        )
        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#6B7280'),
            spaceAfter=15
        )
        h2_style = ParagraphStyle(
            'ReportH2',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=14,
            leading=18,
            textColor=colors.HexColor('#365CF5'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'ReportBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#374151')
        )
        disclaimer_style = ParagraphStyle(
            'ReportDisclaimer',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=8,
            leading=11,
            textColor=colors.HexColor('#6B7280')
        )
        quote_style = ParagraphStyle(
            'ReportQuote',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#4B5563')
        )

        elements = []

        # Header Title
        elements.append(Paragraph("RuleRadar AI — Document Intelligence Report", title_style))
        elements.append(Paragraph(f"Analysis performed on {datetime.now(timezone.utc).strftime('%B %d, %Y at %H:%M UTC')}", subtitle_style))
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E5E7EB'), spaceAfter=15))

        # Metadata Table
        meta_data = [
            [Paragraph("<b>Filename:</b>", body_style), Paragraph(str(doc_data.get("filename", "")), body_style),
             Paragraph("<b>Document Type:</b>", body_style), Paragraph(str(doc_data.get("file_type", "")).upper(), body_style)],
            [Paragraph("<b>Page Count:</b>", body_style), Paragraph(str(doc_data.get("page_count", 1)), body_style),
             Paragraph("<b>Category:</b>", body_style), Paragraph(str(doc_data.get("category", "")).replace("_", " ").title(), body_style)],
            [Paragraph("<b>Total Clauses:</b>", body_style), Paragraph(str(len(doc_data.get("clauses", []))), body_style),
             Paragraph("<b>Risks Flagged:</b>", body_style), Paragraph(str(len(doc_data.get("risk_findings", []))), body_style)]
        ]
        meta_table = Table(meta_data, colWidths=[90, 170, 90, 180])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(meta_table)
        elements.append(Spacer(1, 15))

        # Executive Summary
        elements.append(Paragraph("Executive Summary", h2_style))
        elements.append(Paragraph(doc_data.get("summary") or "Document parsed successfully.", body_style))
        elements.append(Spacer(1, 12))

        # Potential Risk Findings
        risks = doc_data.get("risk_findings", [])
        if risks:
            elements.append(Paragraph(f"Potential Risk Review Indicators ({len(risks)})", h2_style))
            risk_table_rows = [
                [
                    Paragraph("<b>Level</b>", body_style),
                    Paragraph("<b>Issue & Rule</b>", body_style),
                    Paragraph("<b>Verbatim Evidence & Suggested Question</b>", body_style)
                ]
            ]
            for r in risks:
                level = r.get("risk_level", "Review")
                badge_color = "#EF4444" if level == "High attention" else ("#F59E0B" if level == "Needs review" else "#3B82F6")
                level_p = Paragraph(f"<font color='{badge_color}'><b>{level}</b></font>", body_style)
                
                issue_text = f"<b>{r.get('rule_name', '')}</b><br/>{r.get('explanation', '')}"
                issue_p = Paragraph(issue_text, body_style)
                
                evid_text = f"<i>\"{r.get('evidence_quote', '')}\"</i><br/><br/><b>Question to ask:</b> {r.get('suggested_question', '')}"
                evid_p = Paragraph(evid_text, quote_style)
                
                risk_table_rows.append([level_p, issue_p, evid_p])

            risk_table = Table(risk_table_rows, colWidths=[95, 185, 250])
            risk_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
                ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
                ('PADDING', (0, 0), (-1, -1), 6),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))
            elements.append(risk_table)
            elements.append(Spacer(1, 15))

        # Key Contractual Clauses
        clauses = doc_data.get("clauses", [])
        high_clauses = [c for c in clauses if c.get("importance_score", 0.0) >= 5.0]
        if not high_clauses:
            high_clauses = clauses[:6]

        if high_clauses:
            elements.append(Paragraph(f"Key Important Clauses ({len(high_clauses)})", h2_style))
            clause_rows = [
                [
                    Paragraph("<b>Section / Page</b>", body_style),
                    Paragraph("<b>Category & Score</b>", body_style),
                    Paragraph("<b>Original Provision</b>", body_style)
                ]
            ]
            for c in high_clauses[:12]:
                sec = f"{c.get('clause_number') or 'Clause'}<br/>Page {c.get('page_start', 1)}"
                sec_p = Paragraph(sec, body_style)
                
                cat = f"<b>{c.get('category', '')}</b><br/>Score: {c.get('importance_score', 0)}/10 ({c.get('importance_category', '')})"
                cat_p = Paragraph(cat, body_style)
                
                text = c.get('original_text', '')
                if len(text) > 300:
                    text = text[:300] + "..."
                text_p = Paragraph(text, body_style)
                clause_rows.append([sec_p, cat_p, text_p])

            clause_table = Table(clause_rows, colWidths=[80, 160, 290])
            clause_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
                ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
                ('PADDING', (0, 0), (-1, -1), 6),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))
            elements.append(clause_table)
            elements.append(Spacer(1, 15))

        # Deadlines & Obligations
        deadlines = doc_data.get("deadlines", [])
        if deadlines:
            elements.append(Paragraph(f"Extracted Deadlines & Notice Windows ({len(deadlines)})", h2_style))
            for d in deadlines[:8]:
                d_text = f"• <b>{d.get('date_or_period')}</b> ({d.get('deadline_type', 'Deadline')}, Page {d.get('page_number', 1)}): {d.get('context', '')}"
                elements.append(Paragraph(d_text, body_style))
            elements.append(Spacer(1, 12))

        # Legal Disclaimer
        elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#CBD5E1'), spaceAfter=10))
        elements.append(Paragraph(cls.DISCLAIMER_TEXT, disclaimer_style))

        pdf.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def generate_docx_report(cls, doc_data: Dict[str, Any]) -> bytes:
        """
        Generate an editable Microsoft Word (.docx) report using python-docx.
        """
        doc = DocxDoc()

        # Document Header
        title = doc.add_heading("RuleRadar AI — Document Intelligence Report", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT

        sub = doc.add_paragraph(f"Analysis performed: {datetime.now(timezone.utc).strftime('%B %d, %Y at %H:%M UTC')}")
        sub.style.font.color.rgb = RGBColor(107, 114, 128)

        # Metadata
        doc.add_heading("Document Overview", level=1)
        p = doc.add_paragraph()
        p.add_run(f"Filename: {doc_data.get('filename', '')}\n").bold = True
        p.add_run(f"Type: {str(doc_data.get('file_type', '')).upper()} | Pages: {doc_data.get('page_count', 1)} | Category: {doc_data.get('category', '').title()}\n")
        p.add_run(f"Total Clauses Analyzed: {len(doc_data.get('clauses', []))} | Potential Risks Flagged: {len(doc_data.get('risk_findings', []))}\n")

        # Summary
        doc.add_heading("Executive Summary", level=2)
        doc.add_paragraph(doc_data.get("summary") or "Analysis completed.")

        # Risks
        risks = doc_data.get("risk_findings", [])
        if risks:
            doc.add_heading("Potential Risk Review Indicators", level=2)
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Risk Level"
            hdr_cells[1].text = "Issue & Explanation"
            hdr_cells[2].text = "Evidence & Suggested Question"

            for r in risks:
                row_cells = table.add_row().cells
                row_cells[0].text = r.get("risk_level", "")
                row_cells[1].text = f"{r.get('rule_name', '')}\n{r.get('explanation', '')}"
                row_cells[2].text = f"\"{r.get('evidence_quote', '')}\"\n\nQuestion: {r.get('suggested_question', '')}"

        # Clauses
        clauses = doc_data.get("clauses", [])
        if clauses:
            doc.add_heading("Important Clauses Extracted", level=2)
            for c in clauses[:10]:
                p = doc.add_paragraph()
                p.add_run(f"[{c.get('clause_number') or 'Clause'} - {c.get('heading') or 'Section'}] (Page {c.get('page_start', 1)})\n").bold = True
                p.add_run(f"Category: {c.get('category')} | Importance: {c.get('importance_score')}/10\n")
                p.add_run(f"{c.get('original_text', '')}\n")

        # Disclaimer
        doc.add_heading("Notice & Disclaimer", level=2)
        disc_p = doc.add_paragraph(cls.DISCLAIMER_TEXT)
        disc_p.style.font.italic = True

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def generate_json_report(cls, doc_data: Dict[str, Any]) -> str:
        """Export clean structured JSON serialization."""
        return json.dumps(doc_data, indent=2, default=str)
